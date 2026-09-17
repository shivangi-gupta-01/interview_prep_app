"""
Extracts plain text from an uploaded resume file (.pdf, .docx, .txt).
"""
import io
from pathlib import Path

from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
from docx import Document

from app.config import get_settings

settings = get_settings()


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise HTTPException(
            status_code=422,
            detail="Couldn't extract text from this PDF (it may be a scanned image). "
                   "Try a text-based PDF or paste the resume as .txt.",
        )
    return text


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(p for p in parts if p.strip())
    if not text.strip():
        raise HTTPException(status_code=422, detail="Couldn't extract text from this .docx file.")
    return text


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


async def extract_resume_text(file: UploadFile) -> str:
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in settings.allowed_resume_types:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{suffix}'. Allowed: {settings.allowed_resume_types}",
        )
    data = await file.read()
    size_mb = len(data) / (1024 * 1024)
    if size_mb > settings.max_upload_mb:
        raise HTTPException(status_code=413, detail=f"File too large ({size_mb:.1f} MB).")

    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".docx":
        return _extract_docx(data)
    return _extract_txt(data)
